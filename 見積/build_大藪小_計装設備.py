# -*- coding: utf-8 -*-
"""大藪小学校屋内運動場空調設備設置工事　計装設備工事　数量表＋見積
   図面：M-16 機器姿図・計装系統図／M-17 1階平面図／M-18 2階平面図／M-19 屋外平面図／M-20 室外機廻り平面詳細図
   拾い：平面図の経路に書かれた本数〈h×3.b×6〉等 × 区間長（M-17/18/19 は 1/150、M-20 は 1/50、A2判で座標から実測）。
         室内機の立下り 3m、SR-1 引込 9.6m、建物立下り→室外機置場 9m を加え、ケーブルは余長 ×1.1。
   記号：a EM-CEES1.25-2C 集中リモコン／b EM-CEE1.25-2C 個別リモコン／c EM-CEES1.25-7C 自立運転切替SW
         d EM-CEES1.25-2C 遠隔監視／e EM-CE2-3C 遠隔監視アダプタ電源／f EM-CEES1.25-6C 自立信号線
         g 付属ケーブル（メーカー支給）／h EM-CEE1.25-2C 室内機～室外機連絡／i 付属電源ケーブル（別途工事）
   単価：複合単価（積算実務マニュアル2026）。共巻き・天井内転がしは「二重天井内…」（ころがし）の単価。★は登録なし・概算。
"""
import json, base64, math, openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

def jsround(x): return math.floor(x + 0.5)
def sig(v):
    step = max(10, 10 ** (int(math.floor(math.log10(v))) + 1 - 3))
    return int(math.floor(v / step + 0.5)) * step

K150 = 0.05292      # 1/150・A2：1pt = 52.92mm
K50  = 0.00552      # M-20 1/50 を 230dpi で読んだ 1px = 5.52mm
K2F  = 0.04466      # M-18 表示画像 1px = 44.66mm
SEG = []            # (図面, 区間, 長さm, {記号:本数})
def seg(dw, name, m, c): SEG.append((dw, name, round(m, 1), c))
# ---- M-17 1階（大走りに沿って冷媒管共巻き）----
for n, pt, c in [("北 GHP3-2→3-1", 43, {"b":1,"h":1}), ("北 3-1→2-1", 260, {"h":1,"b":2}), ("北 2-2→2-1", 63, {"b":1,"h":1}),
                 ("北 2-1→1-1", 255, {"h":2,"b":4}), ("北 1-2→1-1", 57, {"b":1,"h":1}), ("北 1-1→北東角", 113, {"h":3,"b":6}),
                 ("東 北東角→SR-1分岐", 78, {"h":3,"b":6}), ("東 SR-1分岐→南東角", 409, {"c":3,"h":3,"a":2,"b":3}),
                 ("南東角→立下り", 42, {"c":3,"h":3,"a":2,"b":3}),
                 ("南 GHP3-4→3-3", 43, {"b":1,"h":1}), ("南 3-3→2-3", 262, {"h":3,"b":1}), ("南 2-4→2-3", 58, {"b":1,"h":1}),
                 ("南 2-3→1-3", 255, {"h":5,"b":2}), ("南 1-4→1-3", 56, {"b":1,"h":1}), ("南 1-3→立下り", 69, {"h":7,"b":3})]:
    seg("M-17", n, pt * K150, c)
seg("M-17", "室内機立下り 3m×12台", 36, {"b":1,"h":1})
seg("M-17", "SR-1引込（★コア・PB経由）9.6m", 9.6, {"a":2,"b":3,"c":3})
# ---- M-18 2階（天井内転がし）----
for n, px, c in [("GHP4-1→縦系統", 137, {"b":1,"h":1}), ("縦系統→SR-2分岐", 207, {"b":1,"h":1}), ("SR-2分岐→4-2/4-3分岐", 98, {"b":2,"h":2}),
                 ("分岐→GHP4-2", 70, {"b":1,"h":1}), ("分岐→GHP4-3", 108, {"b":1,"h":1}), ("→外壁（A通り3）", 235, {"h":1})]:
    seg("M-18", n, px * K2F, c)
seg("M-18", "SR-2分岐・盤立下り", 4.8, {"b":1})
seg("M-18", "室内機立下り 3m×3台", 9, {"b":1,"h":1})
seg("M-18", "2F→1F 外壁立下り", 5, {"h":1})
# ---- M-19/20 建物立下り→室外機置場 ----
seg("M-19", "建物立下り→室外機置場", 9, {"a":2,"c":3,"h":4})
for n, px, c in [("置場 入口→GHP-1", 390, {"a":2,"h":4}), ("GHP-1→GHP-2", 320, {"a":2,"d":1,"h":3}), ("GHP-2→GHP-3", 320, {"a":2,"d":1,"h":2}),
                 ("GHP-3→GHP-4", 410, {"a":1,"d":1,"h":1}), ("入口→自立ユニット（共巻き）", 1315, {"c":3})]:
    seg("M-20", n, px * K50, c)
seg("M-20", "機器取合い 0.5m", 0.5, {"a":2,"h":4,"d":3,"c":3})
seg("M-20", "自立ユニット→GHP-1～4", (110 + 300 + 620 + 940) * K50 + 4, {"f":1})
seg("M-20", "GHP-1→遠隔監視アダプタ", 0.8, {"d":1,"e":1})
seg("M-20", "アダプタ電源 立上り", 2, {"e":1})

tot = {}
for _, _, m, c in SEG:
    for k, q in c.items(): tot[k] = tot.get(k, 0) + m * q
L = {k: round(v * 1.1) for k, v in tot.items()}     # 余長10%
CEE2 = L["b"] + L["h"]; CEES2 = L["a"] + L["d"]

# ---- 見積行 ----
rows = []
def cat(n): rows.append({"type": "cat", "name": n})
def it(name, spec, qty, unit, price, pl=0, note=""):
    r = {"name": name, "spec": spec, "qty": qty, "unit": unit, "price": int(price), "note": note}
    if pl: r["pl"] = int(round(pl))
    rows.append(r); return qty * int(price)
QTY = []   # 数量表用（分類, 品名, 仕様, 数量, 単位, 根拠）
def q(cat_, name, spec, n, unit, why): QTY.append((cat_, name, spec, n, unit, why))

sub = 0
cat("01　計装配線")
for nm, sp, n, pr, lab, why in [
    ("ケーブル", "ＥＭ－ＣＥＥ１．２５sq－２Ｃ", CEE2, 690, 338, f"b 個別リモコン {L['b']}m＋h 室内機～室外機 {L['h']}m"),
    ("ケーブル", "ＥＭ－ＣＥＥＳ１．２５sq－２Ｃ", CEES2, 830, 338, f"a 集中リモコン {L['a']}m＋d 遠隔監視 {L['d']}m"),
    ("ケーブル", "ＥＭ－ＣＥＥＳ１．２５sq－７Ｃ", L["c"], 1_660, 677, "c 室外機～自立運転切替SW"),
    ("ケーブル", "ＥＭ－ＣＥＥＳ１．２５sq－６Ｃ", L["f"], 1_440, 564, "f 自立信号線"),
    ("ケーブル", "ＥＭ－ＣＥ２sq－３Ｃ", L["e"], 1_040, 479, "e 遠隔監視アダプタ電源（管内）")]:
    sub += it(nm, sp, n, "ｍ", pr, lab); q("計装配線", nm, sp, n, "ｍ", why)
cat("02　電線管・ボックス")
for nm, sp, n, unit, pr, lab, note, why in [
    ("厚鋼電線管", "Ｇ２８ 屋外露出", 10, "ｍ", 6_320, 3_497, "", "室外機置場 a・d・h（GP-28）"),
    ("厚鋼電線管", "Ｇ３６ 屋外露出", 15, "ｍ", 7_640, 4_202, "", "室外機置場 f・g 自立ユニット～GHP-1～4"),
    ("厚鋼電線管", "Ｇ２２ 屋外露出", 2, "ｍ", 4_840, 2_707, "", "d・e 遠隔監視アダプタ廻り"),
    ("薄鋼電線管", "Ｅ６０ 屋内露出", 6, "ｍ", 8_500, 4_500, "★", "SR-1 引込 a×2・b×3・c×3"),
    ("薄鋼電線管", "Ｅ２５ 屋内露出", 3, "ｍ", 4_030, 2_369, "", "SR-2 立下り"),
    ("プルボックス", "２５０×２５０×２５０ ＳＵＳ製ＷＰ", 1, "個", 38_000, 10_575, "★", "SR-1 引込 屋外側"),
    ("プルボックス", "２５０×２５０×２５０ 錆止塗装", 1, "個", 20_000, 10_575, "", "SR-1 引込 屋内側"),
    ("コア抜き", "φ１００程度", 1, "か所", 24_200, 16_430, "", "★印 SR-1 引込")]:
    sub += it(nm, sp, n, unit, pr, lab, note); q("電線管・ボックス", nm, sp, n, unit, why)
cat("03　空調制御盤")
for nm, sp, pr, lab, why in [
    ("空調制御盤 ＳＲ－１", "自立形 Ｗ６００×Ｄ２００×Ｈ１０００ 架台Ｈ７００", 480_000, 42_300, "鋼板製焼付塗装・鍵No.200、運転リモコン3・ON/OFFコントローラー2・停電時運転盤（自立運転切換SW）組込"),
    ("空調制御盤 ＳＲ－２", "自立形 Ｗ６００×Ｄ２００×Ｈ１０００ 架台Ｈ７００", 260_000, 28_200, "運転リモコン1 組込")]:
    sub += it(nm, sp, 1, "面", pr, lab, "★メーカー見積で差替"); q("空調制御盤", nm, sp, 1, "面", why)
cat("04　試験調整")
sub += it("計装試験調整", "通信確認・自立運転動作確認共", 1, "式", 56_400, 56_400); q("試験調整", "計装試験調整", "", 1, "式", "2人工")
ZATSU = sig(sub * 0.03)
sub += it("雑材消耗品", "", 1, "式", ZATSU, 0, "上記計の3%")

cat("経　費")
labor = sum(r["qty"] * r.get("pl", 0) for r in rows if "qty" in r)
WELFARE, KEIHI = 16.5, 12.0
w_raw = jsround(labor * WELFARE / 100); w_amt = w_raw // 100 * 100
k_raw = jsround(sub * KEIHI / 100)
TARGET = (sub + w_amt + k_raw) // 1000 * 1000
rows.append({"name": "法定福利費", "welfare": WELFARE, "adj": w_amt - w_raw})
rows.append({"name": "諸経費", "rate": KEIHI, "adj": (TARGET - sub - w_amt) - k_raw})

data = {"header": {"name": "大藪小学校屋内運動場空調設備設置工事　計装設備工事", "client": "", "honorific": "御中",
                   "date": "2026-10-03", "staff": "河口", "no": "261005"},
        "place": "", "validity": "発行日より1ヶ月", "remarks": "",
        "notes": ["付属ケーブル（g）・付属電源ケーブル（i）・電源工事は含みません"],
        "taxMode": "ex", "taxRate": 10, "rows": rows}
root = "/home/user/kd-mitsumori"
json.dump(data, open(root + "/見積/見積_大藪小_計装設備.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
json.dump(data, open(root + "/q/oyabu-keisou.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
url = "https://kd-mitsumori.pages.dev/#import=" + base64.urlsafe_b64encode(payload.encode()).decode().rstrip("=")
open(root + "/見積/取込リンク_大藪小_計装設備.txt", "w").write(url + "\nhttps://kd-mitsumori.pages.dev/#q=oyabu-keisou\n")

# ---- 数量表（Excel）----
wb = openpyxl.Workbook(); thin = Side(style="thin"); B = Border(top=thin, bottom=thin, left=thin, right=thin)
hd = PatternFill("solid", fgColor="DDEBF7"); F = Font(name="BIZ UDPゴシック", size=10)
def table(ws, title, head, data, widths):
    ws["A1"] = title; ws["A1"].font = Font(name="BIZ UDPゴシック", size=14, bold=True)
    ws["A2"] = "大藪小学校屋内運動場空調設備設置工事　計装設備工事（M-16～M-20 より拾い出し）"; ws["A2"].font = F
    for j, h in enumerate(head, 1):
        c = ws.cell(4, j, h); c.font = Font(name="BIZ UDPゴシック", size=10, bold=True); c.fill = hd; c.border = B
        c.alignment = Alignment(horizontal="center")
    for i, r in enumerate(data, 5):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, v); c.font = F; c.border = B
    for j, w in enumerate(widths, 1): ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
ws = wb.active; ws.title = "数量表"
table(ws, "計装設備工事　数量表", ["分類", "品名", "仕様", "数量", "単位", "拾い出し根拠"], QTY, [16, 18, 34, 8, 6, 60])
ws2 = wb.create_sheet("拾い出し明細")
KEYS = list("abcdefh")
det = [(dw, n, m) + tuple(c.get(k, "") for k in KEYS) + (round(sum(m * v for v in c.values()), 1),) for dw, n, m, c in SEG]
det.append(("", "計（実長）", "") + tuple(round(tot.get(k, 0), 1) for k in KEYS) + ("",))
det.append(("", "計 ×1.1（余長）→見積数量", "") + tuple(L.get(k, "") for k in KEYS) + ("",))
table(ws2, "計装配線　拾い出し明細（区間長×本数）", ["図面", "区間", "区間長m"] + [f"{k}本数" for k in KEYS] + ["延長m"], det,
      [8, 30, 9] + [7] * 7 + [9])
ws2["A3"] = "a CEES2C集中／b CEE2C個別／c CEES7C自立SW／d CEES2C遠隔監視／e CE2-3C／f CEES6C／h CEE2C連絡　※本数は経路の〈 〉表記による"; ws2["A3"].font = F
wb.save(root + "/見積/数量表_大藪小_計装設備.xlsx")

for r in rows:
    if r.get("type") == "cat": print("【" + r["name"] + "】"); continue
    if "qty" in r: print(f"  {(r['name']+' '+r['spec'])[:34]:36}{r['qty']:>6}{r['unit']:<3}{r['price']:>9,}{r['qty']*r['price']:>11,} {r['note']}")
k_amt = TARGET - sub - w_amt; tax = jsround(TARGET * 0.1)
print(f"小計 {sub:,}／法定福利費 {w_amt:,}／諸経費 {k_amt:,}／計（税抜）{TARGET:,}／税込 {TARGET+tax:,}  人工 {labor/28200:.1f}")
print(L)
