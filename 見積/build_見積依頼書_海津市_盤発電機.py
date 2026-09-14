# -*- coding: utf-8 -*-
"""見積依頼書（商社向け）：各小学校体育館空調設備設置工事 電気設備工事のうち
   動力分電盤（PAC-M／PAC-1）と非常用発電機のみ。
   仕様は図面 E-03 分電盤結線図／E-04 発電機仕様書（山田建築事務所 '26.09）から転記。
   発電機の容量は図面「60kVA以上」→ 100kVA で依頼する。
"""
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "見積依頼書"
F = "ＭＳ Ｐゴシック"
thin = Side(style="thin", color="000000"); med = Side(style="medium", color="000000")
box  = Border(left=thin, right=thin, top=thin, bottom=thin)
head = PatternFill("solid", fgColor="EFEFEF")

ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins.left = ws.page_margins.right = 0.4
ws.page_margins.top = ws.page_margins.bottom = 0.5

W = [4, 12, 17, 50, 6, 5, 11, 12, 10, 22]   # A..J
for i, w in enumerate(W, 1): ws.column_dimensions[get_column_letter(i)].width = w

def put(cell, v, size=10, bold=False, al="left", wrap=False, border=False, fill=False):
    c = ws[cell]; c.value = v
    c.font = Font(name=F, size=size, bold=bold)
    c.alignment = Alignment(horizontal=al, vertical="center", wrap_text=wrap)
    if border: c.border = box
    if fill: c.fill = head
    return c

# ---------------- 表題 ----------------
ws.merge_cells("A1:J1"); put("A1", "御 見 積 依 頼 書", 18, True, "center"); ws.row_dimensions[1].height = 28
ws.merge_cells("G2:J2"); put("G2", "令和８年９月１４日", 10, False, "right")
ws.merge_cells("G3:J3"); put("G3", "依頼No. 260916-A", 10, False, "right")

ws.merge_cells("A5:C5"); put("A5", "", 13, True)
for col in "ABC": ws[col + "5"].border = Border(bottom=thin)
ws.row_dimensions[5].height = 24
put("D5", "御中", 12, True, "left")
ws.merge_cells("A6:D6"); put("A6", "（ 商社名 ）", 9)

ws.merge_cells("G5:J5"); put("G5", "河口電機株式会社", 12, True, "right")
ws.merge_cells("G6:J6"); put("G6", "〒500-8285　岐阜市南鶉６丁目４０－３", 9, False, "right")
ws.merge_cells("G7:J7"); put("G7", "TEL 058-275-4141　FAX 058-275-4133", 9, False, "right")
ws.merge_cells("G8:J8"); put("G8", "担当：河口", 10, False, "right")

ws.merge_cells("A10:J10")
put("A10", "拝啓　時下ますますご清栄のこととお慶び申し上げます。下記の件につきまして、"
           "御見積をいただきたくお願い申し上げます。", 10, False, "left", wrap=True)
ws.row_dimensions[10].height = 26

# ---------------- 件名欄 ----------------
r = 12
info = [("工 事 名 称", "各小学校体育館空調設備設置工事　電気設備工事"),
        ("工 事 場 所", "海津市地内（海西小学校・石津小学校・下多度小学校）"),
        ("発  注  者", "海津市"),
        ("設　　　計", "株式会社 山田建築事務所"),
        ("依 頼 範 囲", "下記【依頼品目】のとおり（動力分電盤・非常用発電機のみ）"),
        ("見積提出期限", "令和８年９月２５日（金）まで"),
        ("納 入 時 期", "別途協議"),
        ("受 渡 場 所", "各小学校 現場搬入"),
        ("そ  の  他", "消費税抜き・現場搬入渡しでお願いします。納期・製作期間も併記願います。")]
for k, v in info:
    ws.merge_cells(f"A{r}:B{r}"); put(f"A{r}", k, 10, True, "center", border=True, fill=True)
    ws.merge_cells(f"C{r}:J{r}"); put(f"C{r}", v, 10, False, "left", border=True)
    for col in "ABCDEFGHIJ": ws[f"{col}{r}"].border = box
    ws.row_dimensions[r].height = 17
    r += 1

# ---------------- 依頼品目 ----------------
r += 1
ws.merge_cells(f"A{r}:J{r}"); put(f"A{r}", "【 依 頼 品 目 】", 12, True); ws.row_dimensions[r].height = 22
r += 1
ws.print_title_rows = f"{r}:{r}"          # 表の見出し行を各ページに繰返し
HDR = ["No", "学　校", "品　　名", "仕　　　　　様", "数量", "単位", "単　価", "金　額", "納　期", "該 当 図 面"]
for i, h in enumerate(HDR, 1):
    c = ws.cell(row=r, column=i, value=h)
    c.font = Font(name=F, size=9, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = box; c.fill = head
ws.row_dimensions[r].height = 22
r += 1

PAC_M2 = ("ＳＵＳ製ＷＰ 壁掛型　3φ3W 200V\n"
          "主幹：MCCB 3P400AF/300AT\n"
          "分岐：MCCB 3P225AF/150AT×1（ＰＡＣ－１へ）\n"
          "　　　ELCB 3P60AF/60AT×2（空調室外機）\n"
          "端子台・銘板共")
PAC_M3 = ("ＳＵＳ製ＷＰ 壁掛型　3φ3W 200V\n"
          "主幹：MCCB 3P400AF/300AT\n"
          "分岐：MCCB 3P225AF/150AT×1（ＰＡＣ－１へ）\n"
          "　　　ELCB 3P60AF/60AT×3（空調室外機）\n"
          "端子台・銘板共")
PAC_1a = ("ＳＵＳ製ＷＰ 壁掛型　3φ3W 200V\n"
          "商用主幹：MCCB 3P225AF/150AT\n"
          "発電機側：ＤＣＳ手動切替スイッチ 3P100A（切替操作扉付）\n"
          "分岐：ELCB 3P60AF/60AT×2（空調室外機）\n"
          "表示灯・補助継電器・変圧器 AC220/DC6V・端子台共")
PAC_1b = ("ＳＵＳ製ＷＰ 壁掛型　3φ3W 200V\n"
          "商用主幹：MCCB 3P225AF/150AT ＋ ELCB 3P60AF/60AT×1\n"
          "発電機側：ＤＣＳ手動切替スイッチ 3P100A（切替操作扉付）\n"
          "　　　　　ELCB 3P60AF/60AT×2（空調室外機）\n"
          "表示灯・補助継電器・変圧器 AC220/DC6V・端子台共")
GEN = ("ディーゼルエンジン発電機　１００ｋＶＡ／210V・60Hz・三相3線・力率80％\n"
       "屋外可搬形　※詳細は下記【非常用発電機 仕様】のとおり。据付・試験運転調整共")
GEN_SPEC = [
 ("容　　量", "１００ｋＶＡ　※図面 Ｅ－０４ は『６０ｋＶＡ以上』／本依頼は１００ｋＶＡ"),
 ("電　　圧", "２１０Ｖ"),
 ("電　　流", "メーカー標準"),
 ("周 波 数", "６０Ｈｚ"),
 ("相　　数", "三相３線"),
 ("力　　率", "８０％"),
 ("励磁方式", "ブラシレス"),
 ("エンジン", "ディーゼル／直接噴射式／ラジエータ冷却／強制潤滑／電気始動（形式・出力・"
             "回転数・気筒数・セルモーター容量・バッテリー容量はメーカー標準）"),
 ("使用燃料", "ＪＩＳ２号軽油（燃料満タンで納入）"),
 ("燃料ﾀﾝｸ", "８時間以上連続運転可能容量"),
 ("そ の 他", "①一般業務用　②手動運転（起動・停止）　③屋外可搬形　④保護装置付　⑤充電器搭載\n"
             "⑥外形寸法は参考とする　⑦箱型出力端子カバー付　⑧排気フランジ・テールパイプ付\n"
             "⑨リアドア・燃料タンクキャップ（もしくはカバー）鍵付\n"
             "⑩「防災用」表示付き保護シートカバー（耐候性防水仕様 厚み０.４７以上）を取付のこと"),
]

ITEMS = [
 ("海西小学校",   "動力分電盤【ＰＡＣ－Ｍ】", PAC_M2, 1, "面", "Ｅ－０３ 分電盤結線図（海西小）",   "設置位置 Ｅ－０５"),
 ("海西小学校",   "動力分電盤【ＰＡＣ－１】", PAC_1a, 1, "面", "Ｅ－０３ 分電盤結線図（海西小）",   "設置位置 Ｅ－０５"),
 ("海西小学校",   "非常用発電機",           GEN,    1, "機", "Ｅ－０４ 発電機仕様書（海西小）",   "設置位置 Ｅ－０２"),
 ("石津小学校",   "動力分電盤【ＰＡＣ－Ｍ】", PAC_M3, 1, "面", "Ｅ－０３ 分電盤結線図（石津小）",   "設置位置 Ｅ－０５"),
 ("石津小学校",   "動力分電盤【ＰＡＣ－１】", PAC_1b, 1, "面", "Ｅ－０３ 分電盤結線図（石津小）",   "設置位置 Ｅ－０５"),
 ("石津小学校",   "非常用発電機",           GEN,    1, "機", "Ｅ－０４ 発電機仕様書（石津小）",   "設置位置 Ｅ－０２"),
 ("下多度小学校", "動力分電盤【ＰＡＣ－Ｍ】", PAC_M3, 1, "面", "Ｅ－０３ 分電盤結線図（下多度小）", "設置位置 Ｅ－０５"),
 ("下多度小学校", "動力分電盤【ＰＡＣ－１】", PAC_1b, 1, "面", "Ｅ－０３ 分電盤結線図（下多度小）", "設置位置 Ｅ－０５"),
 ("下多度小学校", "非常用発電機",           GEN,    1, "機", "Ｅ－０４ 発電機仕様書（下多度小）", "設置位置 Ｅ－０２"),
]
for i, (sch, nm, sp, q, un, dwg, rm) in enumerate(ITEMS, 1):
    vals = [i, sch, nm, sp, q, un, None, None, None, dwg + "\n" + rm]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font = Font(name=F, size=9)
        al = "center" if j in (1, 5, 6) else "left"
        c.alignment = Alignment(horizontal=al, vertical="center", wrap_text=(j in (2, 3, 4, 10)))
        c.border = box
        if j in (7, 8): c.number_format = "#,##0"
    ws.row_dimensions[r].height = 12 * (sp.count("\n") + 1) + 3
    r += 1

# 合計欄
ws.merge_cells(f"A{r}:F{r}"); put(f"A{r}", "合　　計（消費税抜き）", 10, True, "right", border=True)
for col in "ABCDEFGHIJ": ws[f"{col}{r}"].border = box
ws[f"H{r}"].number_format = "#,##0"
ws.row_dimensions[r].height = 22
r += 1

# ---------------- 非常用発電機 仕様 ----------------
r += 1
ws.row_breaks.append(openpyxl.worksheet.pagebreak.Break(id=r - 1))  # 仕様欄から改ページ
ws.merge_cells(f"A{r}:J{r}")
put(f"A{r}", "【 非常用発電機 仕様 】　（図面 Ｅ－０４「発電機仕様書」より　※容量のみ変更）", 12, True)
ws.row_dimensions[r].height = 22
r += 1
for k, v in GEN_SPEC:
    ws.merge_cells(f"A{r}:B{r}"); put(f"A{r}", k, 9.5, True, "center", border=True, fill=True)
    ws.merge_cells(f"C{r}:J{r}"); put(f"C{r}", v, 9.5, False, "left", wrap=True, border=True)
    for col in "ABCDEFGHIJ": ws[f"{col}{r}"].border = box
    ws.row_dimensions[r].height = 12 * (v.count("\n") + 1) + 4
    r += 1

# ---------------- 注記 ----------------
r += 1
ws.merge_cells(f"A{r}:J{r}"); put(f"A{r}", "【 ご 依 頼 に あ た っ て 】", 12, True); ws.row_dimensions[r].height = 22
r += 1
NOTES = [
 "１．非常用発電機の容量は、図面 Ｅ－０４「発電機仕様書」に『６０ｋＶＡ以上』とありますが、"
 "本依頼では　１００ｋＶＡ　で御見積ください。他の項目は同図のとおりです。",
 "２．動力分電盤（ＰＡＣ－Ｍ／ＰＡＣ－１）は、いずれも図面 Ｅ－０３「分電盤結線図」の各校欄によります。"
 "外箱はＳＵＳ製・防水（ＷＰ）仕様の壁掛型です。",
 "３．既設盤（Ｌ－体育館／体育館Ｌ－１）の２回路取替・増設は当社施工のため、本依頼には含みません。",
 "４．発電機は屋外の発電機置場に据付けます（基礎は建築工事）。搬入・据付・試験運転調整の範囲を"
 "明記のうえ御見積ください。",
 "５．単価・金額は消費税抜きでお願いします。あわせて標準製作期間（納期）をご記入ください。",
 "６．図面（Ｅ－０２／Ｅ－０３／Ｅ－０４／Ｅ－０５）を添付します。ご不明な点はご照会ください。",
]
for n in NOTES:
    ws.merge_cells(f"A{r}:J{r}")
    put(f"A{r}", n, 9.5, False, "left", wrap=True)
    ws.row_dimensions[r].height = 15 * (1 + len(n) // 90)
    r += 1

r += 1
ws.merge_cells(f"A{r}:J{r}"); put(f"A{r}", "以　上", 10, False, "right")

for _r in range(1, ws.max_row + 1):
    if ws.row_dimensions[_r].height is None: ws.row_dimensions[_r].height = 13

out = "/home/user/kd-mitsumori/見積/見積依頼書_海津市_分電盤・非常用発電機.xlsx"
wb.save(out); print("保存", out)
